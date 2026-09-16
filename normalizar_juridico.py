"""Normaliza el Excel del Consultorio Jurídico ("CONTROL DE USUARIOS A-C.xlsx")
al modelo legal_consulting.

El Excel trae DOS hojas de grano distinto (no una sola tabla):
  - "CONTROL DE USUARIOS"  -> hecho consulta        (atención al público)
  - "PROCESOS 2025C"       -> hecho proceso_judicial (litigios en curso)

MODO ACUMULAR (piloto de automatización, sep 2026): igual que en Diplomados, los
catálogos y estudiante_consultorio ya NO usan un id numérico propio -- cada fila de los
CSV usa directamente su llave natural (nombre normalizado, código de período, etc.).
generar_sql_carga.py arma con eso un UPSERT (INSERT ... ON CONFLICT ... DO UPDATE), así
funciona igual de bien en la primera carga que acumulando sobre datos existentes.

Excepción importante: proceso_judicial NO tiene llave natural en el Excel de origen (no
trae número de expediente/radicado) -- por eso sigue usando el modo viejo (se genera
aparte, en un archivo de carga MANUAL, ver generar_sql_carga.py). Solo "consulta" entra
al modo acumular / a la automatización de Drive.

ADVERTENCIAS CON NIVEL (mismo acuerdo que Diplomados): cada fila de advertencias_calidad
trae un nivel "critico" o "info".
  - "critico": falta un campo requerido, o un dato no tiene el tipo que le corresponde
    (fecha inválida). La fila se omite. Un proceso automático debe DETENERSE antes de
    tocar CORE si hay alguna advertencia "critico" en la corrida.
  - "info": el dato sí tiene el tipo correcto pero algo es raro (año fuera de rango,
    nombre con coincidencia ambigua). No bloquea la carga, solo queda anotado.

No depende de columnas ni catálogos de Zarigüeya/Tickets.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

DEFAULT_EXCEL = Path(r"C:\Users\angie_vera\Downloads\CONTROL DE USUARIOS A-C.xlsx")
OUT_DIR = Path(__file__).resolve().parent

SHEET_CONSULTAS = "CONTROL DE USUARIOS"
SHEET_PROCESOS = "PROCESOS 2025C"

EMPTY = {"", "-", "NONE", "N/A"}
TICKET_EMPTY = EMPTY | {"NO REGISTRA"}

# Typos de origen detectados en TIPO (hoja PROCESOS) -> nombre canónico.
# Se compara sobre el valor ya "folded" (mayúsculas, sin tildes, espacios colapsados).
TIPO_PROCESO_TYPOS = {
    "EJECUTVO SINGULAR": "EJECUTIVO SINGULAR",
    "EJECUTIVO DE ALIMENTOAS": "EJECUTIVO DE ALIMENTOS",
    "PROCESO DICIPLINARIO": "PROCESO DISCIPLINARIO",
    "DECLARTIVO DE FAMILIA": "DECLARATIVO DE FAMILIA",
}

REQUIRED_CONSULTAS = [
    "FECHA", "HORA", "TICKET", "# SEMESTRE", "MOTIVO DE CONSULTA",
    "ESTUDIANTE ASIGNADO", "MODALIDAD", "ESTADO",
]
REQUIRED_PROCESOS = [
    "LINEA", "PODER DE SUSTITUCION", "ESTUDIANTE ASIGNADO", "ESTADO", "TIPO", "ETAPA",
]


# ---------------------------------------------------------------------------
# Helpers de limpieza (mismo patrón que el pipeline de Diplomados)
# ---------------------------------------------------------------------------

def fold(value: object) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).upper()


def is_empty(value: object, empties: set[str] = EMPTY) -> bool:
    if value is None:
        return True
    return fold(value) in empties


def norm_header(value: object) -> str | None:
    if value is None:
        return None
    return re.sub(r"\s+", " ", str(value).strip()).upper()


def write_csv(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


# ---------------------------------------------------------------------------
# Lectura + validación de hojas
# ---------------------------------------------------------------------------

def find_header_row(raw: list[tuple], required: list[str], sheet_name: str) -> tuple[int, dict[str, int]]:
    """Busca, entre las primeras filas, la que contenga todos los headers requeridos."""
    required_norm = {norm_header(r) for r in required}
    for i, row in enumerate(raw[:5]):
        headers = {norm_header(c): j for j, c in enumerate(row) if c is not None}
        if required_norm <= set(headers.keys()):
            return i, headers
    raise ValueError(
        f"No encuentro los encabezados esperados en la hoja {sheet_name!r}: {required}"
    )


def tokens(name: str) -> list[str]:
    return [t for t in fold(name).split(" ") if t]


def _prefix_match(a: str, b: str) -> bool:
    if a == b:
        return True
    return len(a) >= 4 and len(b) >= 4 and (a.startswith(b) or b.startswith(a))


def is_name_subset(short_tokens: list[str], full_tokens: list[str]) -> bool:
    """True si cada token del nombre corto calza (exacto o por prefijo >=4) con
    algún token del nombre completo. Sirve para unir 'PABLO GAMA' con
    'PABLO ANDRES GAMA VILLAMIL' sin forzar coincidencias débiles."""
    return all(any(_prefix_match(st, ft) for ft in full_tokens) for st in short_tokens)


def normalizar(excel: Path, out_dir: Path | None = None) -> dict[str, int]:
    dest = out_dir or OUT_DIR
    dest.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(excel, data_only=True, read_only=True)

    sheet_names = {norm_header(n): n for n in wb.sheetnames}
    if norm_header(SHEET_CONSULTAS) not in sheet_names or norm_header(SHEET_PROCESOS) not in sheet_names:
        wb.close()
        raise ValueError(
            f"El Excel debe tener las hojas {SHEET_CONSULTAS!r} y {SHEET_PROCESOS!r}. "
            f"Encontradas: {wb.sheetnames}"
        )

    ws1 = wb[sheet_names[norm_header(SHEET_CONSULTAS)]]
    raw1 = list(ws1.iter_rows(values_only=True))
    ws2 = wb[sheet_names[norm_header(SHEET_PROCESOS)]]
    raw2 = list(ws2.iter_rows(values_only=True))
    wb.close()

    h1_idx, col1 = find_header_row(raw1, REQUIRED_CONSULTAS, SHEET_CONSULTAS)
    h2_idx, col2 = find_header_row(raw2, REQUIRED_PROCESOS, SHEET_PROCESOS)
    data1 = [r for r in raw1[h1_idx + 1:] if any(c is not None for c in r)]
    data2 = [r for r in raw2[h2_idx + 1:] if any(c is not None for c in r)]

    def c1(row: tuple, name: str) -> object:
        idx = col1[norm_header(name)]
        return row[idx] if idx < len(row) else None

    def c2(row: tuple, name: str) -> object:
        idx = col2[norm_header(name)]
        return row[idx] if idx < len(row) else None

    # -- catálogos: solo se guarda la llave natural + su(s) atributo(s) descriptivo(s),
    #    sin id propio (Postgres lo asigna vía UPSERT, ver generar_sql_carga.py) ----------
    periodos: "OrderedDict[str, tuple[str, int]]" = OrderedDict()   # codigo -> (letra_origen, anio)
    areas: "OrderedDict[str, None]" = OrderedDict()                  # nombre (folded)
    modalidades: "OrderedDict[str, None]" = OrderedDict()            # nombre (folded)
    estados: "OrderedDict[str, None]" = OrderedDict()                # codigo==nombre (folded)
    tipos: "OrderedDict[str, str]" = OrderedDict()                   # canonico -> origen
    estudiantes_origen: "OrderedDict[str, str]" = OrderedDict()      # normalizado -> origen (primera vez visto)

    consultas: list[list[object]] = []
    procesos: list[list[object]] = []
    advertencias: list[list[object]] = []

    def estudiante_full(nombre_crudo: object) -> str | None:
        if is_empty(nombre_crudo):
            return None
        norm = fold(nombre_crudo)
        estudiantes_origen.setdefault(norm, str(nombre_crudo).strip())
        return norm

    def estudiante_short(nombre_crudo: object, fila_ref: str) -> str | None:
        if is_empty(nombre_crudo):
            return None
        norm = fold(nombre_crudo)
        if norm in estudiantes_origen:
            return norm
        short_tok = tokens(norm)
        candidatos = [
            existing_norm
            for existing_norm in estudiantes_origen.keys()
            if is_name_subset(short_tok, tokens(existing_norm))
        ]
        if len(candidatos) == 1:
            return candidatos[0]
        if len(candidatos) > 1:
            advertencias.append([
                fila_ref, "ESTUDIANTE ASIGNADO", nombre_crudo,
                f"coincidencia ambigua con {candidatos}; se crea registro nuevo, revisar a mano",
                "info",
            ])
        estudiantes_origen.setdefault(norm, str(nombre_crudo).strip())
        return norm

    # -- hoja 1: CONTROL DE USUARIOS -> consulta --------------------------------
    for i, row in enumerate(data1, start=h1_idx + 2):  # +2: 1-based y salta header
        fila_ref = f"{SHEET_CONSULTAS}!fila{i}"
        fecha_val = c1(row, "FECHA")
        letra = c1(row, "# SEMESTRE")
        motivo = c1(row, "MOTIVO DE CONSULTA")
        estudiante = c1(row, "ESTUDIANTE ASIGNADO")
        modalidad = c1(row, "MODALIDAD")
        estado_raw = c1(row, "ESTADO")

        faltantes = [
            campo for campo, val in (
                ("FECHA", fecha_val), ("# SEMESTRE", letra), ("MOTIVO DE CONSULTA", motivo),
                ("ESTUDIANTE ASIGNADO", estudiante), ("MODALIDAD", modalidad), ("ESTADO", estado_raw),
            ) if is_empty(val)
        ]
        if faltantes:
            advertencias.append([fila_ref, ",".join(faltantes), "", "campo requerido vacío; fila omitida de consulta", "critico"])
            continue

        if isinstance(fecha_val, datetime):
            fecha = fecha_val.date()
        elif isinstance(fecha_val, date):
            fecha = fecha_val
        else:
            advertencias.append([fila_ref, "FECHA", fecha_val, "formato de fecha inesperado (no es fecha de Excel); fila omitida", "critico"])
            continue
        if fecha.year < 2020:
            advertencias.append([fila_ref, "FECHA", fecha.isoformat(), "año fuera de rango (posible typo, ej. 2006 en vez de 2026); revisar Excel de origen", "info"])

        letra_txt = str(letra).strip().upper()
        periodo_codigo = f"{fecha.year}{letra_txt}"
        periodos.setdefault(periodo_codigo, (letra_txt, fecha.year))

        area_nombre = fold(motivo)
        areas.setdefault(area_nombre, None)
        estudiante_norm = estudiante_full(estudiante)
        modalidad_nombre = fold(modalidad)
        modalidades.setdefault(modalidad_nombre, None)

        estado_folded = fold(estado_raw)
        estados.setdefault(estado_folded, None)

        ticket = c1(row, "TICKET")
        ticket_txt = "" if is_empty(ticket, TICKET_EMPTY) else str(ticket).strip()
        hora_txt = "" if is_empty(c1(row, "HORA")) else str(c1(row, "HORA")).strip()

        consultas.append([
            ticket_txt, fecha.isoformat(), hora_txt, periodo_codigo, area_nombre,
            estudiante_norm, modalidad_nombre, estado_folded,
        ])

    # -- hoja 2: PROCESOS 2025C -> proceso_judicial (sigue en modo MANUAL, ver LEEME) ---
    for i, row in enumerate(data2, start=h2_idx + 2):
        fila_ref = f"{SHEET_PROCESOS}!fila{i}"
        linea = c2(row, "LINEA")
        poder = c2(row, "PODER DE SUSTITUCION")
        estudiante = c2(row, "ESTUDIANTE ASIGNADO")
        estado_raw = c2(row, "ESTADO")
        tipo_raw = c2(row, "TIPO")
        etapa_raw = c2(row, "ETAPA")  # opcional: 1 nulo observado en origen

        faltantes = [
            campo for campo, val in (
                ("LINEA", linea), ("ESTUDIANTE ASIGNADO", estudiante),
                ("ESTADO", estado_raw), ("TIPO", tipo_raw),
            ) if is_empty(val)
        ]
        if faltantes:
            advertencias.append([fila_ref, ",".join(faltantes), "", "campo requerido vacío; fila omitida de proceso_judicial", "critico"])
            continue

        area_nombre = fold(linea)
        areas.setdefault(area_nombre, None)
        estudiante_norm = estudiante_short(estudiante, fila_ref)

        estado_txt = str(estado_raw).strip()
        if "/" in estado_txt:
            estado_base, observacion = estado_txt.split("/", 1)
        else:
            estado_base, observacion = estado_txt, ""
        estado_folded = fold(estado_base)
        estados.setdefault(estado_folded, None)
        observacion = observacion.strip()

        tipo_folded = fold(tipo_raw)
        tipo_canonico = TIPO_PROCESO_TYPOS.get(tipo_folded, tipo_folded)
        tipos.setdefault(tipo_canonico, str(tipo_raw).strip())

        poder_bool = None
        if not is_empty(poder):
            poder_bool = fold(poder) in {"SI", "SÍ", "S", "TRUE", "1"}

        etapa_txt = "" if is_empty(etapa_raw) else str(etapa_raw).strip()

        procesos.append([
            area_nombre, poder_bool, estudiante_norm, estado_folded, observacion,
            tipo_canonico, etapa_txt, "",  # etapa_fecha: no se infiere, ver 01_legal_consulting_schema.sql
        ])

    # -- volcado a CSV / xlsx (sin columna "id": la llave natural va en cada fila) ------
    periodo_rows = [[codigo, letra, anio] for codigo, (letra, anio) in periodos.items()]
    area_rows = [[n] for n in areas.keys()]
    modalidad_rows = [[n] for n in modalidades.keys()]
    estado_rows = [[n] for n in estados.keys()]
    tipo_rows = [[n, origen] for n, origen in tipos.items()]
    estudiante_rows = [[origen, norm] for norm, origen in estudiantes_origen.items()]

    tables = {
        "periodo_academico": (["codigo", "letra_origen", "anio"], periodo_rows),
        "area_derecho": (["nombre"], area_rows),
        "modalidad_atencion": (["nombre"], modalidad_rows),
        "estado_caso": (["codigo"], estado_rows),
        "tipo_proceso": (["nombre_normalizado", "nombre_origen"], tipo_rows),
        "estudiante_consultorio": (["nombre_origen", "nombre_normalizado"], estudiante_rows),
        "consulta": (
            ["ticket_externo", "fecha", "hora_texto", "periodo_codigo", "area_derecho_nombre",
             "estudiante_normalizado", "modalidad_nombre", "estado_codigo"],
            consultas,
        ),
        "proceso_judicial": (
            ["area_derecho_nombre", "poder_sustitucion", "estudiante_normalizado", "estado_codigo",
             "estado_observacion", "tipo_proceso_nombre", "etapa_texto", "etapa_fecha"],
            procesos,
        ),
    }

    out_xlsx = Workbook()
    first = True
    for name, (hdrs, rows) in tables.items():
        write_csv(dest / f"{name}.csv", hdrs, rows)
        sheet = out_xlsx.active if first else out_xlsx.create_sheet(name)
        if first:
            sheet.title = name
            first = False
        sheet.append(hdrs)
        for row in rows:
            sheet.append(list(row))

    n_criticas = sum(1 for a in advertencias if a[-1] == "critico")
    if advertencias:
        write_csv(dest / "advertencias_calidad.csv", ["fila", "campo", "valor", "detalle", "nivel"], advertencias)

    xlsx_path = dest / "juridico_normalizado.xlsx"
    try:
        out_xlsx.save(xlsx_path)
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        xlsx_path = dest / f"juridico_normalizado_{stamp}.xlsx"
        out_xlsx.save(xlsx_path)
        print(
            "juridico_normalizado.xlsx está abierto (Excel o Cursor). "
            f"Se guardó copia en {xlsx_path.name}. Cierra el original para la próxima."
        )

    print(f"consultas={len(consultas)} procesos={len(procesos)} (procesos sigue en modo MANUAL, no se automatiza)")
    print(f"periodos={len(periodo_rows)} areas={len(area_rows)} modalidades={len(modalidad_rows)}")
    print(f"estados={len(estado_rows)} tipos_proceso={len(tipo_rows)} estudiantes={len(estudiante_rows)}")
    print(f"advertencias={len(advertencias)} (criticas={n_criticas})" + (" -> ver advertencias_calidad.csv" if advertencias else ""))
    print(f"xlsx={xlsx_path}")

    return {
        "consultas": len(consultas),
        "procesos": len(procesos),
        "estudiantes": len(estudiante_rows),
        "advertencias": len(advertencias),
        "advertencias_criticas": n_criticas,
    }


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Normaliza el Excel de Jurídico al modelo legal_consulting.")
    parser.add_argument(
        "excel",
        nargs="?",
        default=str(DEFAULT_EXCEL),
        help="Ruta del xlsx de Jurídico (default: Descargas/CONTROL DE USUARIOS A-C.xlsx)",
    )
    args = parser.parse_args()
    path = Path(args.excel)
    if not path.exists():
        raise SystemExit(f"No existe el Excel: {path}")
    normalizar(path, OUT_DIR)


if __name__ == "__main__":
    main()
